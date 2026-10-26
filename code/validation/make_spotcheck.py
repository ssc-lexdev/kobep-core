#!/usr/bin/env python3
"""Draw the stratified spot-check sample (C11) and write the review workbook.

Every other validation in KOBEP is machine against machine - parser against the
DART API, totals against published aggregates. This one puts a person in front
of the original filing. The sample is drawn from the release files themselves,
so what is checked is what is published.

Strata and sizes
----------------
filings       160  4 receipt-year bands x {standard, short}, 20 each; up to 3
                   amendments per stratum (amendments carry the pre/post table
                   trap, pitfall F4)
transactions  200  by flow_kind, split as evenly as possible before/after 2016
events        100  class_at_entry x entry_left_truncated

The draw is seeded, so the same release yields the same sample. Reviewers mark
each field OK / WRONG / CANT_TELL; `score_spotcheck.py` turns the marks into
error rates with Wilson 95 % intervals for the Technical Validation section.
"""
from __future__ import annotations

import argparse
import os

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

SEED = 20261001
VIEWER = "https://dart.fss.or.kr/dsaf001/main.do?rcpNo={}"
MARKS = '"OK,WRONG,CANT_TELL"'


def band(year: pd.Series) -> pd.Series:
    return pd.cut(year.astype(int), [2010, 2015, 2019, 2022, 2025],
                  labels=["2011-15", "2016-19", "2020-22", "2023-25"]).astype(str)


def take(df: pd.DataFrame, n: int, rng: np.random.Generator) -> pd.DataFrame:
    if len(df) <= n:
        return df
    return df.iloc[np.sort(rng.choice(len(df), n, replace=False))]


def sample_filings(d: str, rng) -> pd.DataFrame:
    p = pd.read_csv(os.path.join(d, "kobep_position_ledger.csv"), dtype=str)
    ent = pd.read_csv(os.path.join(d, "kobep_blockholder_entity.csv"), dtype=str,
                      usecols=["reporter_id", "name"])
    firm = (pd.read_csv(os.path.join(d, "kobep_firm_master.csv"), dtype=str,
                        usecols=["corp_code", "corp_name_ko"])
            .drop_duplicates("corp_code", keep="last"))
    p["stratum"] = band(p["receipt_date"].str[:4]) + " " + p["form_kind"]
    picks = []
    for _, g in p.groupby("stratum", sort=True):
        amended = take(g[g["is_amendment"] == "1"], 3, rng)
        picks.append(amended)
        picks.append(take(g[g["is_amendment"] == "0"], 20 - len(amended), rng))
    s = pd.concat(picks).merge(ent, on="reporter_id", how="left").merge(
        firm, on="corp_code", how="left")
    s["dart_url"] = s["rcept_no"].map(VIEWER.format)
    cols = ["stratum", "rcept_no", "dart_url", "corp_name_ko", "name", "receipt_date",
            "is_amendment", "obligation_date", "report_type_code", "relation_to_issuer_code",
            "prev_ratio_stockequiv", "curr_ratio_stockequiv", "holding_purpose_code",
            "holding_purpose_code_derived", "holding_purpose_basis", "total_shares_issued"]
    return s[cols].rename(columns={"name": "filer_name"})


def sample_transactions(d: str, rng) -> pd.DataFrame:
    t = pd.read_csv(os.path.join(d, "kobep_transaction_ledger.csv"), dtype=str)
    quota = {"TRADE": 80, "OTHER": 40, "NON_CASH": 20, "PERIMETER": 20,
             "UNKNOWN": 15, "LENDING": 15, "DERIVATIVE": 10}
    picks = []
    for kind, n in quota.items():
        g = t[t["flow_kind"] == kind]
        early = take(g[g["in_core_window"] == "0"], n // 2, rng)
        picks.append(early)
        picks.append(take(g[g["in_core_window"] == "1"], n - len(early), rng))
    s = pd.concat(picks)
    s["stratum"] = s["flow_kind"] + np.where(s["in_core_window"] == "1", " 2016-25", " 2011-15")
    s["dart_url"] = s["rcept_no"].map(VIEWER.format)
    cols = ["stratum", "rcept_no", "row_no", "dart_url", "party_name", "change_date",
            "method_ko", "security_ko", "shares_before", "shares_change", "shares_after",
            "unit_price"]
    return s[cols]


def sample_events(d: str, rng) -> pd.DataFrame:
    e = pd.read_csv(os.path.join(d, "kobep_event_ledger.csv"), dtype=str)
    ent = pd.read_csv(os.path.join(d, "kobep_blockholder_entity.csv"), dtype=str,
                      usecols=["reporter_id", "name"])
    idx = pd.read_csv(os.path.join(d, "kobep_filing_entity.csv"), dtype=str).merge(
        pd.read_csv(os.path.join(d, "kobep_filing_index.csv"), dtype=str,
                    usecols=["rcept_no", "corp_code", "receipt_date"]), on="rcept_no")
    quota = {("OUTSIDER", "0"): 30, ("OUTSIDER", "1"): 15, ("UNCERTAIN", "0"): 15,
             ("UNCERTAIN", "1"): 10, ("INSIDER", "0"): 15, ("INSIDER", "1"): 15}
    picks = [take(e[(e["class_at_entry"] == c) & (e["entry_left_truncated"] == tr)], n, rng)
             .assign(stratum=f"{c} truncated={tr}") for (c, tr), n in quota.items()]
    s = pd.concat(picks).merge(ent, on="reporter_id", how="left")
    filings = (idx.sort_values("receipt_date").groupby(["corp_code", "reporter_id"])["rcept_no"]
                  .apply(lambda x: " ".join(x)).rename("pair_filings"))
    s = s.merge(filings, on=["corp_code", "reporter_id"], how="left")
    s["entry_dart_url"] = s["entry_rcept_no"].map(VIEWER.format)
    cols = ["stratum", "event_id", "corp_code", "name", "entry_rcept_no", "entry_dart_url",
            "entry_obligation_date", "entry_receipt_date", "entry_ratio", "prev_ratio",
            "class_at_entry", "entry_left_truncated", "is_first_ever", "exit_type",
            "exit_date", "last_observed_ratio", "pair_filings"]
    return s[cols].rename(columns={"name": "filer_name"})


CHECKS = {
    "filings": ["obligation_date", "report_type_code", "relation_to_issuer_code",
                "prev_ratio_stockequiv", "curr_ratio_stockequiv", "holding_purpose_code",
                "holding_purpose_code_derived", "total_shares_issued", "filer_name"],
    "transactions": ["party_name", "change_date", "method_ko", "shares_change",
                     "shares_after", "unit_price"],
    "events": ["entry_obligation_date", "entry_ratio", "class_at_entry",
               "entry_left_truncated", "exit_type", "exit_date"],
}


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--release", default=r"C:\KOBEP\release\KOBEP-Core_v1.1")
    ap.add_argument("--out", required=True)
    args = ap.parse_args()
    d = os.path.join(args.release, "data")
    rng = np.random.default_rng(SEED)
    sheets = {"filings": sample_filings(d, rng),
              "transactions": sample_transactions(d, rng),
              "events": sample_events(d, rng)}
    for name, df in sheets.items():
        for c in CHECKS[name]:
            df[f"chk_{c}"] = ""
        df["reviewer_note"] = ""
    tmp = args.out + ".tmp.xlsx"
    with pd.ExcelWriter(tmp, engine="openpyxl") as xw:
        pd.DataFrame({"how to review": [
            "Open dart_url (or entry_dart_url) and compare each value with the original filing.",
            "Mark every chk_ column OK, WRONG or CANT_TELL (drop-down). A blank value is OK if the filing is blank there too.",
            "Ratios: compare with the post-amendment table on amended filings (pitfall F4).",
            "events: entry_left_truncated = 1 means an earlier holding is evident (unparsed earlier filing, or prior ratio >= 5 %). pair_filings lists the pair's filings in order.",
            "Write anything odd in reviewer_note. Do not edit the value columns.",
            f"Sample seed {SEED}; regenerate with make_spotcheck.py."]}).to_excel(
            xw, sheet_name="README", index=False)
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
    wb = load_workbook(tmp)
    for name, df in sheets.items():
        ws = wb[name]
        dv = DataValidation(type="list", formula1=MARKS, allow_blank=True)
        ws.add_data_validation(dv)
        fill = PatternFill("solid", fgColor="FFF2CC")
        for j, col in enumerate(df.columns, 1):
            cell = ws.cell(1, j)
            cell.font = Font(bold=True)
            if col.startswith("chk_") or col == "reviewer_note":
                cell.fill = fill
                if col.startswith("chk_"):
                    letter = cell.column_letter
                    dv.add(f"{letter}2:{letter}{len(df) + 1}")
            if col.endswith("dart_url"):
                for i in range(2, len(df) + 2):
                    c = ws.cell(i, j)
                    c.hyperlink = c.value
                    c.font = Font(color="0563C1", underline="single")
        ws.freeze_panes = "C2"
        for colcells in ws.columns:
            ws.column_dimensions[colcells[0].column_letter].width = min(
                40, max(10, max(len(str(c.value or "")) for c in colcells[:50]) + 2))
    wb["README"].column_dimensions["A"].width = 120
    for row in wb["README"].iter_rows(min_row=2):
        row[0].alignment = Alignment(wrap_text=True)
    wb.save(tmp)
    os.replace(tmp, args.out)
    for name, df in sheets.items():
        print(f"{name:13s} {len(df):4d} items   strata: {df['stratum'].nunique()}")
    print(f"-> {args.out}")


if __name__ == "__main__":
    main()
