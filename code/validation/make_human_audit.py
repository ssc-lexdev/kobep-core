#!/usr/bin/env python3
"""Stage 2 of C11: the human-audit workbook.

Takes the LLM-screened workbook (spotcheck_C11_llm.xlsx) and keeps only the
rows the human must decide: every row the LLM flagged and the seeded 10 % of
all-OK rows. For each row it shows our value, the LLM verdict, the value the
LLM read and its evidence, plus an assistant triage note - a provisional
reading, explicitly not a verdict. The reviewer fills human_<field> with
OK / WRONG / CANT_TELL for every checked field, so the LLM's own accuracy can
be measured on these rows.
"""
from __future__ import annotations

import argparse
import os
import sqlite3

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

FIELDS = {
    "filings": ["obligation_date", "report_type_code", "relation_to_issuer_code",
                "prev_ratio_stockequiv", "curr_ratio_stockequiv", "holding_purpose_code",
                "holding_purpose_code_derived", "total_shares_issued", "filer_name"],
    "transactions": ["party_name", "change_date", "method_ko", "shares_change",
                     "shares_after", "unit_price"],
    "events": ["entry_obligation_date", "entry_ratio", "class_at_entry",
               "entry_left_truncated", "exit_type", "exit_date"],
}
SUSPECT_MERGES = {"한앤코시멘트홀딩스주식회사", "제이에프이스틸"}


def triage(sheet: str, row: pd.Series, reentries: dict) -> str:
    notes = []
    for f in FIELDS[sheet]:
        v = row.get(f"llm_{f}", "")
        if v not in ("WRONG", "CANT_TELL"):
            continue
        orig = str(row.get(f"llm_orig_{f}", ""))
        if sheet == "events" and f in ("exit_type", "exit_date") and reentries.get(row["event_id"], 0):
            notes.append(f"{f}: 재진입 관계(n_reentries={reentries[row['event_id']]}) — 설계상 마지막 하락만 이탈, 정상일 가능성 높음")
        elif sheet == "events" and f == "class_at_entry" and ("목적" in orig or "경영권" in orig):
            notes.append(f"{f}: LLM이 분류(INSIDER/OUTSIDER/UNCERTAIN)와 보유목적을 혼동")
        elif sheet == "events" and f == "entry_left_truncated" and "대표보고자" in str(row.get("llm_evidence", "")):
            notes.append(f"{f}: 대표보고자 변동 신규보고 — 정의상 0, 경제적으로는 기존 보유(정의 보강 후보)")
        elif sheet == "filings" and f == "filer_name":
            if row.get("filer_name") in SUSPECT_MERGES:
                notes.append(f"{f}: 🔴 서로 다른 법인을 한 엔티티로 합친 의심 — 확인 필요")
            else:
                notes.append(f"{f}: 엔티티 정규화명 vs 신고 당시 상호(상호변경 가능)")
        elif sheet == "transactions" and f == "party_name" and v == "CANT_TELL":
            notes.append(f"{f}: 발췌 한계(여러 줄 이름 셀) — 원문에서 전체 이름 확인")
        elif sheet == "filings" and f == "holding_purpose_code_derived":
            notes.append(f"{f}: 파생 규칙 미결 표현 — LLM 판단 {orig}, 규칙 보강 후보")
        else:
            notes.append(f"{f}: 확인 필요 (LLM 원문값 {orig[:30]})")
    if not notes:
        notes.append("LLM 전부 OK — 무작위 10% 감사 표본")
    return " / ".join(notes)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("llm_workbook")
    ap.add_argument("--out", required=True)
    args = ap.parse_args()
    con = sqlite3.connect("file:C:/KOBEP/ledger.db?mode=ro", uri=True)
    reentries = dict(con.execute("select event_id, n_reentries from treatment_event"))
    sheets = pd.read_excel(args.llm_workbook, sheet_name=None, dtype=str)
    out = {}
    for sheet, fields in FIELDS.items():
        df = sheets[sheet].fillna("")
        df = df[df["human_audit"] != ""].copy()
        df.insert(0, "assistant_triage", [triage(sheet, r, reentries) for _, r in df.iterrows()])
        keep = [c for c in df.columns if not c.startswith("chk_") and c != "reviewer_note"]
        df = df[keep]
        for f in fields:
            df[f"human_{f}"] = ""
        df["human_note"] = ""
        out[sheet] = df
    tmp = args.out + ".tmp.xlsx"
    with pd.ExcelWriter(tmp, engine="openpyxl") as xw:
        pd.DataFrame({"how to review": [
            "Rows here are (a) every item the LLM marked WRONG/CANT_TELL and (b) a seeded 10 % of items it marked all-OK.",
            "Open dart_url (or entry_dart_url). For EVERY human_<field> column, enter OK / WRONG / CANT_TELL for OUR value (not the LLM's).",
            "assistant_triage is Claude's provisional reading to speed you up. It is not a verdict - overrule it freely.",
            "llm_<field> = LLM verdict, llm_orig_<field> = value the LLM read, llm_evidence = its quote.",
            "Then run validation/score_spotcheck.py on this file plus the LLM workbook (combined scoring)."]}).to_excel(
            xw, sheet_name="README", index=False)
        for sheet, df in out.items():
            df.to_excel(xw, sheet_name=sheet, index=False)
    wb = load_workbook(tmp)
    for sheet, df in out.items():
        ws = wb[sheet]
        dv = DataValidation(type="list", formula1='"OK,WRONG,CANT_TELL"', allow_blank=True)
        ws.add_data_validation(dv)
        for j, col in enumerate(df.columns, 1):
            cell = ws.cell(1, j)
            cell.font = Font(bold=True)
            if col.startswith("human_"):
                cell.fill = PatternFill("solid", fgColor="FFF2CC")
                if col != "human_note":
                    dv.add(f"{cell.column_letter}2:{cell.column_letter}{len(df) + 1}")
            if col.endswith("dart_url"):
                for i in range(2, len(df) + 2):
                    c = ws.cell(i, j)
                    c.hyperlink = c.value
                    c.font = Font(color="0563C1", underline="single")
        ws.column_dimensions["A"].width = 70
        ws.freeze_panes = "B2"
    wb["README"].column_dimensions["A"].width = 130
    wb.save(tmp)
    os.replace(tmp, args.out)
    for s, df in out.items():
        print(f"{s:13s} {len(df):3d} rows")
    print(f"-> {args.out}")


if __name__ == "__main__":
    main()
