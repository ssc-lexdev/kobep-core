#!/usr/bin/env python3
"""Stage 3 of C11: human check of the rows the human audit did not cover.

Takes the LLM-screened workbook (spotcheck_C11_llm.xlsx) and keeps the rows
with a blank human_audit flag - the rows that so far carry only an LLM verdict.
The LLM verdict, the value it read and its evidence are left out, so the
reviewer checks our value against the original filing blind to the model.
The reviewer fills human_<field> with OK / WRONG / CANT_TELL for every field.
Together with the earlier audit this gives a human verdict on every sampled field.

--ledger c11_field_ledger.csv (2026-10-26) instead keeps every sampled row with
any field the human has not judged yet (human blank in the ledger written by
score_stratified.py). Fields of those rows that already carry a human verdict
are pre-filled with "-" and are not judged again.
"""
from __future__ import annotations

import argparse
import os

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from make_human_audit import FIELDS

KEYS = {"filings": ["rcept_no"], "transactions": ["rcept_no", "row_no"], "events": ["event_id"]}

README = [
    "Rows here are the sampled rows that the earlier human audit did not cover (human_audit blank in spotcheck_C11_llm.xlsx).",
    "The LLM verdicts are deliberately not shown: check OUR value against the original filing, blind to the model.",
    "Open dart_url (events: entry_dart_url, then the pair's later filings for exit_type / exit_date).",
    "For EVERY human_<field> column enter OK / WRONG / CANT_TELL. Use human_note for anything unusual (give the value you read when WRONG).",
    "Field rules are the same as in the README of spotcheck_C11.xlsx (e.g. exit = start of the FINAL run below 5 %, relationships observed up to 2025-12-31).",
    "When done, score with validation/score_stratified.py together with the earlier audit.",
    "A cell pre-filled with '-' already has a human verdict from an earlier sheet: leave it.",
    "",
    "[한국어] LLM 판정은 일부러 숨겼습니다. 원문을 열어 '우리 값'이 맞는지만 판정하세요. 모든 human_ 칸에 OK / WRONG / CANT_TELL을 넣고, WRONG이면 human_note에 원문 값을 적어 주세요. '-' 칸은 이미 판정한 칸이니 그대로 두세요.",
]


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("llm_workbook")
    ap.add_argument("--out", required=True)
    ap.add_argument("--sample", type=int, default=0,
                    help="draw this many rows in total, allocated to sheets in proportion to their "
                         "remaining rows (0 = all remaining rows)")
    ap.add_argument("--seed", type=int, default=20261022)
    ap.add_argument("--ledger", help="c11_field_ledger.csv: keep rows with any field not yet judged by the human")
    args = ap.parse_args()
    sheets = pd.read_excel(args.llm_workbook, sheet_name=None, dtype=str)
    done = {}
    if args.ledger:
        led = pd.read_csv(args.ledger, dtype=str).fillna("")
        open_keys = set(zip(led.sheet[led.human == ""], led.key[led.human == ""]))
        done = {(r.sheet, r.key, r.field) for r in led[led.human != ""].itertuples()}
        remaining = {}
        for s in FIELDS:
            d = sheets[s].fillna("")
            k = d[KEYS[s]].agg("#".join, axis=1)
            remaining[s] = d[[(s, x) in open_keys for x in k]]
        assert sum(len(d) for d in remaining.values()) == len({k for k in open_keys}), "ledger keys not found"
    else:
        remaining = {s: sheets[s].fillna("").query("human_audit == ''") for s in FIELDS}
    total = sum(len(d) for d in remaining.values())
    alloc = {s: len(d) for s, d in remaining.items()}
    if args.sample:
        # Largest-remainder allocation so the sheet quotas sum exactly to --sample.
        raw = {s: args.sample * len(d) / total for s, d in remaining.items()}
        alloc = {s: int(v) for s, v in raw.items()}
        for s in sorted(raw, key=lambda s: raw[s] - alloc[s], reverse=True)[:args.sample - sum(alloc.values())]:
            alloc[s] += 1
    out = {}
    for sheet, fields in FIELDS.items():
        df = remaining[sheet].copy()
        if args.sample:
            df = df.sample(n=alloc[sheet], random_state=args.seed).sort_index()
            df.insert(0, "sample_frame", f"{alloc[sheet]} of {len(remaining[sheet])} remaining rows, seed {args.seed}")
        keep = [c for c in df.columns
                if not c.startswith(("chk_", "llm_")) and c not in ("reviewer_note", "human_audit")]
        df = df[keep]
        keys = df[KEYS[sheet]].agg("#".join, axis=1)
        for f in fields:
            df[f"human_{f}"] = ["-" if (sheet, k, f) in done else "" for k in keys]
        df["human_note"] = ""
        out[sheet] = df
    tmp = args.out + ".tmp.xlsx"
    with pd.ExcelWriter(tmp, engine="openpyxl") as xw:
        pd.DataFrame({"how to review": README}).to_excel(xw, sheet_name="README", index=False)
        for sheet, df in out.items():
            df.to_excel(xw, sheet_name=sheet, index=False)
    wb = load_workbook(tmp)
    for sheet, df in out.items():
        ws = wb[sheet]
        dv = DataValidation(type="list", formula1='"OK,WRONG,CANT_TELL"', allow_blank=True)
        ws.add_data_validation(dv)
        for j, col in enumerate(df.columns, 1):
            cell = ws.cell(1, j)
            cell.font = Font(bold=True)
            if col.startswith("human_"):
                cell.fill = PatternFill("solid", fgColor="FFF2CC")
                ws.column_dimensions[cell.column_letter].width = 14
                if col != "human_note":
                    dv.add(f"{cell.column_letter}2:{cell.column_letter}{len(df) + 1}")
                    for i in range(2, len(df) + 2):
                        if ws.cell(i, j).value == "-":
                            ws.cell(i, j).fill = PatternFill("solid", fgColor="D9D9D9")
                else:
                    ws.column_dimensions[cell.column_letter].width = 40
            if col.endswith("dart_url"):
                for i in range(2, len(df) + 2):
                    c = ws.cell(i, j)
                    c.hyperlink = c.value
                    c.font = Font(color="0563C1", underline="single")
        ws.freeze_panes = "B2"
    wb["README"].column_dimensions["A"].width = 150
    wb.save(tmp)
    os.replace(tmp, args.out)
    total = 0
    for s, df in out.items():
        n = int(sum((df[f"human_{f}"] == "").sum() for f in FIELDS[s]))
        total += n
        print(f"{s:13s} {len(df):3d} rows, {n:,} fields to judge")
    print(f"total {total:,} fields -> {args.out}")


if __name__ == "__main__":
    main()
