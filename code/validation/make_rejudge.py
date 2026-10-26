#!/usr/bin/env python3
"""C11 spot check: re-judgement sheet for fields whose released value changed
after the audit (2026-10-23).

The audit judged the values of an earlier build. Two later changes moved some
of them: the holding-purpose rules (2026-10-05 and 2026-10-23) and the 2025-12-31
observation cutoff for events (2026-10-07). Each changed field is judged twice:

  human_then_b   the audited value under the criterion fixed on 2026-10-23
                 (a derived purpose left blank although the prose can be
                 classified is WRONG). `proposed_then_b` pre-fills the
                 criterion's mechanical answer for blank derived purposes;
                 it is a proposal, the human column decides.
  human_current  the value in the current release.

Events in the sample that the new `entry_rep_change` flag marks are added
under that field, with no audited value, so the flag itself is checked.

Inputs (deliverables/spotcheck): c11_field_ledger.csv; release data folder.
Output: spotcheck_C11_rejudge.xlsx
"""
from __future__ import annotations

import os
import sys

import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

D = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "deliverables", "spotcheck")
REL = sys.argv[1] if len(sys.argv) > 1 else r"C:\KOBEP\release\KOBEP-Core_v1.1\data"


def num(v: str):
    try:
        return float(v)
    except ValueError:
        return v.strip()


def main() -> None:
    led = pd.read_csv(os.path.join(D, "c11_field_ledger.csv"), dtype=str).fillna("")
    pos = pd.read_csv(os.path.join(REL, "kobep_position_ledger.csv"), dtype=str).fillna("")
    ev = pd.read_csv(os.path.join(REL, "kobep_event_ledger.csv"), dtype=str).fillna("")
    tx = pd.read_csv(os.path.join(REL, "kobep_transaction_ledger.csv"), dtype=str).fillna("")
    tx["key"] = tx.rcept_no + "#" + tx.row_no
    cur = {"filings": pos.set_index("rcept_no"), "events": ev.set_index("event_id"),
           "transactions": tx.set_index("key")}

    rows = []
    for r in led.itertuples():
        table = cur[r.sheet]
        if r.key not in table.index or r.field not in table.columns:
            continue  # filer_name is not a position-ledger column
        now = table.at[r.key, r.field]
        if num(now) == num(r.our_value):
            continue
        # The criterion is about prose: a blank that only the form cover fills
        # (basis FORM_COVER) was not classifiable from the prose.
        blank_purpose = (r.field == "holding_purpose_code_derived" and r.our_value == ""
                         and table.at[r.key, "holding_purpose_basis"] == "TEXT_RULE")
        rows.append({"sheet": r.sheet, "key": r.key, "field": r.field,
                     "audited_value": r.our_value, "current_value": now,
                     "human_then": r.human, "human_then_source": r.human_source,
                     "stratum": r.stratum,
                     "proposed_then_b": "WRONG" if blank_purpose and now else "",
                     "evidence": r.claude_evidence,
                     "human_then_b": "", "human_current": "", "human_note": ""})
    sampled = set(led.loc[led.sheet == "events", "key"])
    for eid, e in cur["events"].iterrows():
        if eid in sampled and e["entry_rep_change"] == "1":
            rows.append({"sheet": "events", "key": eid, "field": "entry_rep_change",
                         "audited_value": "(not audited)", "current_value": "1",
                         "human_then": "", "human_then_source": "", "stratum": "",
                         "proposed_then_b": "", "evidence": f"entry filing {e['entry_rcept_no']}",
                         "human_then_b": "n/a", "human_current": "", "human_note": ""})
    df = pd.DataFrame(rows)
    out = os.path.join(D, "spotcheck_C11_rejudge.xlsx")
    df.to_excel(out, index=False, sheet_name="rejudge")
    wb = load_workbook(out)
    ws = wb["rejudge"]
    col = {c.value: c.column_letter for c in ws[1]}
    dv = DataValidation(type="list", formula1='"OK,WRONG,CANT_TELL"', allow_blank=True)
    ws.add_data_validation(dv)
    for c in ("human_then_b", "human_current"):
        dv.add(f"{col[c]}2:{col[c]}{ws.max_row}")
    ws.column_dimensions[col["evidence"]].width = 70
    ws.freeze_panes = "D2"
    wb.save(out)
    print(f"{len(df)} fields -> {out}")
    print(df.groupby(["sheet", "field"]).size().to_string())


if __name__ == "__main__":
    main()
